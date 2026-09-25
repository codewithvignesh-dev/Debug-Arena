import os, re, io, json, time, random, secrets, hmac, functools
import pymysql
from flask import Flask, request, session, redirect, render_template, jsonify, send_file, abort
from werkzeug.security import generate_password_hash, check_password_hash as cph
from openpyxl import load_workbook
from dotenv import load_dotenv
load_dotenv()
E = os.environ.get
app = Flask(__name__)
app.secret_key = E('SECRET_KEY')
app.config.update(SESSION_COOKIE_HTTPONLY=True, SESSION_COOKIE_SAMESITE='Lax')
QN, QT, GR, MK = 10, 90, 3, 10  # questions, seconds per question, grace seconds, marks per question
SL = {'new': 'Not Started', 'run': 'In Progress', 'done': 'Completed', 'dq': 'Disqualified'}
gph = lambda p: generate_password_hash(p, method='pbkdf2:sha256:100000')

def q(sql, a=(), one=False):
    c = pymysql.connect(host=E('DB_HOST', 'localhost'), user=E('DB_USER'), password=E('DB_PASS'),
                        database=E('DB_NAME'), charset='utf8mb4', autocommit=True,
                        cursorclass=pymysql.cursors.DictCursor)
    try:
        with c.cursor() as k:
            n = k.execute(sql, a)
            if sql.lstrip()[:6].upper() != 'SELECT': return n
            return k.fetchone() if one else k.fetchall()
    finally:
        c.close()

def init():
    for s in [
        "CREATE TABLE IF NOT EXISTS users(id INT AUTO_INCREMENT PRIMARY KEY, username VARCHAR(50) UNIQUE, pw VARCHAR(255), role VARCHAR(10))",
        "CREATE TABLE IF NOT EXISTS students(id INT AUTO_INCREMENT PRIMARY KEY, reg VARCHAR(20) UNIQUE, name VARCHAR(80), pw VARCHAR(255), status VARCHAR(8) DEFAULT 'new', st BIGINT DEFAULT 0, et BIGINT DEFAULT 0, viol INT DEFAULT 0, lv BIGINT DEFAULT 0, qst BIGINT DEFAULT 0, mcq TINYINT DEFAULT 0)",
        "CREATE TABLE IF NOT EXISTS sq(id INT AUTO_INCREMENT PRIMARY KEY, sid INT, pos INT, title VARCHAR(60), code TEXT, ans VARCHAR(40), opts TEXT NULL, given VARCHAR(40) NULL, ok TINYINT NULL, at BIGINT DEFAULT 0, INDEX(sid))",
        "CREATE TABLE IF NOT EXISTS alerts(id INT AUTO_INCREMENT PRIMARY KEY, reg VARCHAR(20), msg VARCHAR(120), ts BIGINT)",
        "CREATE TABLE IF NOT EXISTS settings(k VARCHAR(20) PRIMARY KEY, v VARCHAR(20))"]:
        q(s)
    try: q("ALTER TABLE students ADD COLUMN qst BIGINT DEFAULT 0")
    except Exception: pass
    try: q("ALTER TABLE students ADD COLUMN mcq TINYINT DEFAULT 0")
    except Exception: pass
    try: q("ALTER TABLE sq ADD COLUMN opts TEXT NULL")
    except Exception: pass
    q("INSERT IGNORE INTO settings VALUES('ev','wait'),('show','0')")
    q("INSERT INTO users(username,pw,role) VALUES(%s,%s,'admin') ON DUPLICATE KEY UPDATE pw=VALUES(pw), role='admin'", (E('ADMIN_USER'), gph(E('ADMIN_PASS'))))
init()

# ---------- question bank: 100 fixed beginner C++ tracing questions ----------
from question_bank import QBANK

def ev(): return q("SELECT v FROM settings WHERE k='ev'", one=True)['v']
def visible(): return ev() == 'done' and q("SELECT v FROM settings WHERE k='show'", one=True)['v'] == '1'
def sweep():
    q("UPDATE students SET status='done' WHERE status='run' AND (et<%s OR %s)", (int(time.time()), 1 if ev() == 'done' else 0))
def me():
    return q("SELECT * FROM students WHERE id=%s", (session.get('sid'),), True) if session.get('sid') else None
def cur(sid): return q("SELECT id,pos,title,code,ans,opts FROM sq WHERE sid=%s AND given IS NULL ORDER BY pos LIMIT 1", (sid,), True)
def sync(s):
    if s['status'] != 'run': return s
    now = int(time.time())
    while s['status'] == 'run' and now >= s['qst'] + QT + GR:  # question time over -> skip it
        if q("UPDATE students SET qst=%s WHERE id=%s AND qst=%s", (s['qst'] + QT, s['id'], s['qst'])) == 1:
            q("UPDATE sq SET given='' WHERE sid=%s AND given IS NULL ORDER BY pos LIMIT 1", (s['id'],))
        s = q("SELECT * FROM students WHERE id=%s", (s['id'],), True)
        if not cur(s['id']): s['status'] = 'done'
    if s['status'] == 'run' and (now > s['et'] or ev() == 'done' or not cur(s['id'])): s['status'] = 'done'
    if s['status'] == 'done': q("UPDATE students SET status='done' WHERE id=%s AND status='run'", (s['id'],))
    return s
def mcq_options(ans, r):
    """Build 4 shuffled options (one correct) for a numeric-answer question."""
    opts = {ans}
    try:
        base = int(ans)
        deltas = [d for d in range(-6, 7) if d != 0]
        r.shuffle(deltas)
        for d in deltas:
            if len(opts) == 4: break
            cand = str(base + d)
            if cand != ans and base + d >= 0: opts.add(cand)
        tries = 0
        while len(opts) < 4 and tries < 50:
            cand = str(max(0, base + r.randint(-25, 25))); opts.add(cand); tries += 1
    except ValueError:
        for p in ['0', '1', '2', '3', '4', '5', '6', '7', '8', '9', '10']:
            if len(opts) == 4: break
            if p != ans: opts.add(p)
    lst = list(opts)[:4]
    r.shuffle(lst)
    return lst

def start(s):
    now = int(time.time())
    if q("UPDATE students SET status='run',st=%s,et=%s,qst=%s WHERE id=%s AND status='new'", (now, now + QN * (QT + GR), now, s['id'])) != 1: return
    r = random.Random(secrets.randbits(64))
    pick = r.sample(QBANK, QN)  # 10 unique questions out of the 100-question bank, own shuffled order per student
    for i, t in enumerate(pick):
        opts = json.dumps(mcq_options(t['answer'], r)) if s['mcq'] else None
        q("INSERT INTO sq(sid,pos,title,code,ans,opts) VALUES(%s,%s,%s,%s,%s,%s)", (s['id'], i, t['title'], t['code'], t['answer'], opts))

def need(*roles):
    def d(f):
        @functools.wraps(f)
        def w(*a, **k):
            if session.get('role') not in roles: return redirect('/')
            return f(*a, **k)
        return w
    return d

@app.context_processor
def _(): return dict(t=session.setdefault('t', secrets.token_hex(16)))
@app.before_request
def csrf():
    if request.method == 'POST':
        tok = request.form.get('_t') or request.headers.get('X-T') or ''
        if not hmac.compare_digest(str(session.get('t', '')).encode(), tok.encode()): abort(400)
@app.after_request
def hdr(r):
    r.headers.update({'X-Frame-Options': 'DENY', 'X-Content-Type-Options': 'nosniff', 'Cache-Control': 'no-store'}); return r

FAIL = {}
@app.route('/', methods=['GET', 'POST'])
def login():
    err = None
    if request.method == 'POST':
        ip = request.remote_addr; FAIL[ip] = [t for t in FAIL.get(ip, []) if t > time.time() - 300]
        role, u, p = request.form.get('role'), request.form.get('u', '').strip(), request.form.get('p', '')
        if len(FAIL[ip]) >= 10: err = 'Too many attempts. Wait 5 minutes.'
        elif role == 'student':
            s = q("SELECT * FROM students WHERE reg=%s", (u.upper(),), True)
            if s and cph(s['pw'], p):
                session.clear(); session.update(sid=s['id'], t=secrets.token_hex(16)); return redirect('/exam')
        elif role in ('staff', 'admin'):
            w = q("SELECT * FROM users WHERE username=%s AND role=%s", (u, role), True)
            if w and cph(w['pw'], p):
                session.clear(); session.update(role=role, t=secrets.token_hex(16)); return redirect('/dash')
        if not err: FAIL[ip].append(time.time()); err = 'Invalid register number / username or password.'
    return render_template('login.html', err=err)

@app.route('/logout')
def logout(): session.clear(); return redirect('/')

# ---------- student ----------
@app.route('/exam')
def exam():
    s = me()
    if not s: return redirect('/')
    s = sync(s)
    if s['status'] == 'new':
        if ev() != 'live': return render_template('exam.html', s=s, mode='wait')
        start(s)
    return render_template('exam.html', s=s, mode='run')

@app.route('/api/q')
def api_q():
    s = me()
    if not s: abort(401)
    s = sync(s)
    if s['status'] != 'run': return jsonify(st=s['status'])
    row = cur(s['id'])
    if not row:
        q("UPDATE students SET status='done' WHERE id=%s", (s['id'],)); return jsonify(st='done')
    return jsonify(st='run', n=row['pos'] + 1, total=QN, title=row['title'], code=row['code'],
                   opts=json.loads(row['opts']) if row['opts'] else None, left=max(0, int(s['qst'] + QT - time.time())))

@app.route('/api/answer', methods=['POST'])
def api_a():
    s = me()
    if not s: abort(401)
    s = sync(s)
    if s['status'] == 'run':
        j = request.get_json(silent=True) or {}; a = str(j.get('a', '')).strip()[:40]; now = int(time.time()); row = cur(s['id'])
        if row and j.get('n') == row['pos'] + 1 and (a or now >= s['qst'] + QT - 3):  # blank answer only counts as skip when time is up
            if q("UPDATE sq SET given=%s, ok=%s, at=%s WHERE id=%s AND given IS NULL", (a, int(a == row['ans']) if a else None, now if a else 0, row['id'])) == 1:
                q("UPDATE students SET qst=%s WHERE id=%s", (now, s['id']))
    return api_q()

@app.route('/api/viol', methods=['POST'])
def viol():
    s = me()
    if not s: abort(401)
    s = sync(s)
    if s['status'] != 'run': return jsonify(st=s['status'])
    now = int(time.time())
    if now - s['lv'] < 3: return jsonify(st='run', v=0)
    v = s['viol'] + 1; dq = v >= 3
    q("UPDATE students SET viol=%s, lv=%s, status=%s WHERE id=%s", (v, now, 'dq' if dq else 'run', s['id']))
    if v >= 2:
        q("INSERT INTO alerts(reg,msg,ts) VALUES(%s,%s,%s)", (s['reg'], 'DISQUALIFIED - 3rd tab switch' if dq else 'Warning 2/3 - left the exam tab', now))
    return jsonify(st='dq' if dq else 'run', v=v)

# ---------- staff / admin ----------
def rows():
    R = q("SELECT s.reg,s.name,s.status,s.viol,s.st,COALESCE(SUM(x.ok=1),0) c,COALESCE(SUM(x.ok=0),0) w,COALESCE(MAX(x.at),0) la "
          "FROM students s LEFT JOIN sq x ON x.sid=s.id AND x.given IS NOT NULL GROUP BY s.id,s.reg,s.name,s.status,s.viol,s.st")
    for r in R:
        r['c'], r['w'] = int(r['c']), int(r['w']); r['un'] = QN - r['c'] - r['w']
        r['tt'] = int(r['la'] - r['st']) if r['la'] else 0; r['sl'] = SL[r['status']]; r['score'] = r['c'] * MK
    R.sort(key=lambda r: (-r['c'], r['tt'] or 10**9))
    for i, r in enumerate(R): r['rank'] = i + 1
    return R

@app.route('/dash')
@need('staff', 'admin')
def dash():
    st = q("SELECT id,username FROM users WHERE role='staff' ORDER BY username") if session['role'] == 'admin' else []
    return render_template('dash.html', role=session['role'], m=request.args.get('m', ''), staff=st)

@app.route('/api/stats')
@need('staff', 'admin')
def stats():
    sweep()
    c = {r['status']: r['n'] for r in q("SELECT status,COUNT(*) n FROM students GROUP BY status")}
    a = q("SELECT COALESCE(SUM(ok=1),0) c, COALESCE(SUM(ok=0),0) w, COALESCE(SUM(given=''),0) k FROM sq WHERE given IS NOT NULL", one=True)
    vis = visible()
    return jsonify(total=sum(c.values()), new=c.get('new', 0), run=c.get('run', 0), done=c.get('done', 0), dq=c.get('dq', 0),
                   correct=int(a['c']), wrong=int(a['w']), skipped=int(a['k']), ev=ev(), rows=rows() if vis else None,
                   alerts=q("SELECT reg,msg,ts FROM alerts ORDER BY id DESC LIMIT 40"))

@app.route('/export.xlsx')
@need('staff', 'admin')
def export():
    sweep()
    if not visible(): abort(403)
    from openpyxl import Workbook
    wb = Workbook(); ws = wb.active; ws.title = 'Results'
    ws.append(['Rank', 'Register No', 'Name', 'Status', f'Marks (out of {QN*MK})', 'Correct', 'Wrong', 'Skipped', 'Time Taken (mm:ss)', 'Tab Warnings'])
    for r in rows():
        ws.append([r['rank'], r['reg'], r['name'], r['sl'], r['score'], r['c'], r['w'], r['un'], f"{r['tt']//60}:{r['tt']%60:02d}" if r['tt'] else '-', r['viol']])
    for col, wd in zip('ABCDEFGHIJ', (7, 16, 26, 14, 18, 9, 9, 9, 20, 13)): ws.column_dimensions[col].width = wd
    b = io.BytesIO(); wb.save(b); b.seek(0)
    return send_file(b, as_attachment=True, download_name='debug_event_results.xlsx',
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

@app.route('/admin/ev', methods=['POST'])
@need('admin')
def a_ev():
    v = request.form.get('v')
    if v in ('wait', 'live', 'done'): q("UPDATE settings SET v=%s WHERE k='ev'", (v,))
    return redirect('/dash')

@app.route('/admin/show', methods=['POST'])
@need('admin')
def a_show():
    q("UPDATE settings SET v=IF(v='1','0','1') WHERE k='show'"); return redirect('/dash?m=Results visibility toggled')

@app.route('/admin/reset', methods=['POST'])
@need('admin')
def a_reset():
    q("DELETE FROM sq"); q("DELETE FROM alerts")
    q("UPDATE students SET status='new',st=0,et=0,viol=0,lv=0"); q("UPDATE settings SET v='wait' WHERE k='ev'"); q("UPDATE settings SET v='0' WHERE k='show'")
    return redirect('/dash?m=Event reset')

REG_RE = re.compile(r'^[0-9A-Z]{6,12}$')  # register no: letters+digits, no spaces (e.g. 2026U437, 25U11C001, 2422K0838)
FIRSTYEAR_RE = re.compile(r'^2026U\d{3}$')  # this batch's first-year reg-no format -> gets MCQ questions instead of typed answers

def import_pairs(pairs):
    """pairs: list of (cell_a, cell_b) from a pasted line or a spreadsheet row, order unknown."""
    n = skipped = 0
    for a, b in pairs:
        a, b = str(a).strip(), str(b).strip()
        if not a or not b: skipped += 1; continue
        if REG_RE.match(a.upper()) and not REG_RE.match(b.upper()): reg, name = a, b
        elif REG_RE.match(b.upper()): reg, name = b, a
        else: skipped += 1; continue
        reg, name = reg.upper(), name[:80]
        pw = name[:3].capitalize() + '@sngc#' + reg[-3:]  # first 3 letters of name + @sngc# + last 3 characters of register no
        mcq = 1 if FIRSTYEAR_RE.match(reg) else 0
        try:
            q("INSERT INTO students(reg,name,pw,mcq) VALUES(%s,%s,%s,%s) ON DUPLICATE KEY UPDATE name=VALUES(name), pw=VALUES(pw), mcq=VALUES(mcq)", (reg, name, gph(pw), mcq))
            n += 1
        except Exception:
            skipped += 1
    return n, skipped

@app.route('/admin/students', methods=['POST'])
@need('admin')
def a_students():
    pairs = []
    for ln in request.form.get('data', '').splitlines():
        p = [x.strip() for x in re.split(r'[,\t]', ln) if x.strip()]
        if len(p) >= 2: pairs.append((p[0], p[1]))
    n, skipped = import_pairs(pairs)
    msg = f'{n} students imported' + (f', {skipped} lines skipped (no valid register number found)' if skipped else '')
    return redirect(f'/dash?m={msg}')

@app.route('/admin/students/upload', methods=['POST'])
@need('admin')
def a_students_xlsx():
    f = request.files.get('file')
    if not f or not f.filename.lower().endswith(('.xlsx', '.xlsm')):
        return redirect('/dash?m=Please upload a .xlsx or .xlsm file')
    try:
        wb = load_workbook(filename=io.BytesIO(f.read()), read_only=True, data_only=True)
        ws = wb.active
        pairs = []
        for row in ws.iter_rows(values_only=True):
            cells = [str(c).strip() for c in row if c is not None and str(c).strip() != '']
            if len(cells) >= 2: pairs.append((cells[0], cells[1]))
    except Exception:
        return redirect('/dash?m=Could not read that Excel file. Make sure it is a valid .xlsx.')
    n, skipped = import_pairs(pairs)
    msg = f'{n} students imported from Excel' + (f', {skipped} rows skipped' if skipped else '')
    return redirect(f'/dash?m={msg}')

@app.route('/api/students')
@need('admin')
def api_students():
    return jsonify(q("SELECT id,reg,name,status,mcq FROM students ORDER BY reg"))

@app.route('/api/student/update', methods=['POST'])
@need('admin')
def api_student_update():
    d = request.get_json(silent=True) or {}
    sid, reg, name, pwd = d.get('id'), str(d.get('reg', '')).strip().upper(), str(d.get('name', '')).strip()[:80], str(d.get('password', '')).strip()
    mcq = 1 if str(d.get('mcq', '0')) in ('1', 'true', 'True') else 0
    if not sid or not REG_RE.match(reg) or not name:
        return jsonify(ok=False, error='Enter a valid register number and name.')
    if pwd and len(pwd) < 4:
        return jsonify(ok=False, error='Password must be at least 4 characters.')
    try:
        if pwd: q("UPDATE students SET reg=%s,name=%s,pw=%s,mcq=%s WHERE id=%s", (reg, name, gph(pwd), mcq, sid))
        else: q("UPDATE students SET reg=%s,name=%s,mcq=%s WHERE id=%s", (reg, name, mcq, sid))
    except pymysql.err.IntegrityError:
        return jsonify(ok=False, error='That register number is already used by another student.')
    return jsonify(ok=True)

@app.route('/api/student/delete', methods=['POST'])
@need('admin')
def api_student_delete():
    sid = (request.get_json(silent=True) or {}).get('id')
    if not sid: return jsonify(ok=False)
    q("DELETE FROM sq WHERE sid=%s", (sid,)); q("DELETE FROM students WHERE id=%s", (sid,))
    return jsonify(ok=True)


@app.route('/admin/staff', methods=['POST'])
@need('admin')
def a_staff():
    u, p = request.form.get('u', '').strip(), request.form.get('p', '')
    if not re.fullmatch(r'[A-Za-z0-9_.-]{3,30}', u) or len(p) < 6: return redirect('/dash?m=Username needs 3-30 letters or numbers, password needs 6+ characters')
    try: q("INSERT INTO users(username,pw,role) VALUES(%s,%s,'staff')", (u, gph(p)))
    except pymysql.err.IntegrityError: return redirect('/dash?m=That username already exists')
    return redirect('/dash?m=Staff account added')

@app.route('/admin/staff/del', methods=['POST'])
@need('admin')
def a_staff_del():
    q("DELETE FROM users WHERE id=%s AND role='staff'", (request.form.get('id'),)); return redirect('/dash?m=Staff account removed')

if __name__ == '__main__': app.run(debug=False)